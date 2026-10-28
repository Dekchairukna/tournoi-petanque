"""Run with: python3 -m unittest discover -s tests"""
import io
import os
import unittest

# Keep all integration-test writes away from the user's database.
os.environ['DATABASE_URL'] = 'sqlite:///:memory:'
os.environ['YAGAMI_SUPERADMIN_PASSWORD'] = 'local-test-owner-password'
from app import app, db
from models import Event, Team, User
from openpyxl import Workbook, load_workbook


class EventImportTests(unittest.TestCase):
    def setUp(self):
        app.config.update(TESTING=True, WTF_CSRF_ENABLED=False)
        self.context = app.app_context()
        self.context.push()
        db.create_all()
        user = User(username='import-test', role='superadmin', password_hash='unused')
        db.session.add(user)
        db.session.commit()
        self.client = app.test_client()
        with self.client.session_transaction() as session:
            session['_user_id'] = str(user.id)
            session['_fresh'] = True

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.context.pop()

    def upload(self, rows, selected=None):
        workbook = Workbook()
        for row in rows:
            workbook.active.append(row)
        stream = io.BytesIO()
        workbook.save(stream)
        stream.seek(0)
        data = {'file': (stream, 'events.xlsx')}
        if selected:
            data['competition_format'] = selected
        response = self.client.post('/event/import', data=data)
        self.assertEqual(response.status_code, 302)

    def test_select_each_format_for_multiple_events(self):
        for fmt in ('swiss', 'round_robin', 'double_knockout', 'knockout'):
            self.upload([['event_name', 'team_name', 'rounds'],
                         [fmt + ' A', 'Team A', 5], [fmt + ' B', 'Team B', 5]], fmt)
            events = Event.query.filter(Event.name.like(fmt + ' %')).all()
            self.assertEqual(len(events), 2)
            self.assertTrue(all(e.competition_format == fmt for e in events))
            self.assertTrue(all(e.rounds == (5 if fmt == 'swiss' else 1) for e in events))
        self.assertEqual(Team.query.count(), 8)

    def test_mixed_formats_and_blank_rows(self):
        self.upload([['event_key', 'event_name', 'team_name', 'competition_format'],
                     ['A', 'A', 'One', None], ['A', 'A', 'Two', 'Round Robin'],
                     ['A', 'A', 'One', None], ['B', 'B', 'Three', 'swiss'],
                     ['C', 'C', 'Four', None]], 'knockout')
        self.assertEqual([(e.name, e.competition_format) for e in Event.query.order_by(Event.id)],
                         [('A', 'round_robin'), ('B', 'swiss'), ('C', 'knockout')])
        self.assertEqual(Team.query.count(), 4)

    def test_old_file_defaults_to_swiss(self):
        self.upload([['event_name', 'team_name'], ['Old', 'One']])
        self.assertEqual(Event.query.one().competition_format, 'swiss')
        self.assertEqual(Event.query.one().rounds, 3)

    def test_invalid_or_conflicting_format_rolls_back_whole_file(self):
        for values in [('typo', 'typo'), ('swiss', 'knockout')]:
            self.upload([['event_name', 'team_name', 'competition_format'],
                         ['Valid', 'One', 'swiss'], ['Bad', 'Two', values[0]],
                         ['Bad', 'Three', values[1]]])
            self.assertEqual(Event.query.count(), 0)
            self.assertEqual(Team.query.count(), 0)

    def test_form_and_template(self):
        response = self.client.get('/')
        self.assertEqual(response.status_code, 200)
        self.assertIn(b'id="bulkCompetitionFormat"', response.data)
        response = self.client.get('/event/import-template')
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(io.BytesIO(response.data)).active
        self.assertIn('competition_format', [cell.value for cell in sheet[1]])


if __name__ == '__main__':
    unittest.main()
